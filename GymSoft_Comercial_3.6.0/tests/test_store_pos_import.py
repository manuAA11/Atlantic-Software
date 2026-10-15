import json
import tempfile
import unittest
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from cloud_database import CloudDatabase, EXCEL_IMPORT_SHEETS
from store_pos_import import POS_BACKUP_SHEETS, validate_pos_import


class StoreImport(unittest.TestCase):
    def database(self):
        order = str(uuid4())
        self.tables = {
            "clients": [dict(id=1, document="P-1", first_name="Persona", active=True)],
            "plans": [dict(id=2, name="Mensual", duration_days=30, price=70000)],
            "store_products": [dict(id=3, name="Coca-Cola", sku="00034", barcode="007701234567890",
                                    description="Botella", category="Bebidas", unit_cost=2000,
                                    sale_price=4000, stock_quantity=10, low_stock_threshold=5, active=True)],
            "store_sales": [dict(id=4, product_id=3, product_sku="00034", product_barcode="007701234567890",
                                 product_name="Coca-Cola", quantity=2, unit_price=4000, unit_cost=2000,
                                 total_amount=8000, order_id=order, status="COMPLETED",
                                 payment_method="EFECTIVO", sold_at="2026-10-15T02:30:00Z")],
            "store_orders": [dict(id=order, sale_number=125, request_id=str(uuid4()),
                                  request_payload={"items": [{"product_id": 3, "quantity": 2}]},
                                  client_id=1, customer_name="Persona", total_amount=8000,
                                  status="COMPLETED", sold_at="2026-10-15T02:30:00Z")],
            "store_payments": [dict(id=str(uuid4()), order_id=order, method="EFECTIVO", amount=8000,
                                    received=10000, change_amount=2000, status="COMPLETED",
                                    created_at="2026-10-15T02:30:00Z")],
            "store_inventory_movements": [dict(id=5, product_id=3, quantity=-2, stock_before=12,
                                               stock_after=10, movement_type="SALE", order_id=order,
                                               created_at="2026-10-15T02:30:00Z")],
            "gym_store_settings": [dict(enabled=True, last_sale_number=125,
                                        updated_at="2026-10-15T02:30:00Z")],
        }
        db = CloudDatabase.__new__(CloudDatabase)
        db.gym_id = "local-test"
        db.timezone = "America/Bogota"
        db._cloud_table_rows = lambda table: deepcopy(self.tables.get(table, []))
        db._pos_backup_snapshot = lambda: deepcopy({table: self.tables.get(table, [])
            for table in (*POS_BACKUP_SHEETS, "store_products", "store_sales")})
        db._now = lambda: datetime(2026, 10, 15, 2, 30, tzinfo=timezone.utc)
        return db

    def test_complete_excel_roundtrip_keeps_distinct_codes_cost_and_full_pos_ledger(self):
        db = self.database()
        with tempfile.TemporaryDirectory() as folder:
            path = db.export_excel(Path(folder) / "pos.xlsx")
            tables = db.prepare_excel_replacement(path)["tables"]
        product = tables["store_products"][0]
        self.assertEqual((product["sku"], product["barcode"]), ("00034", "007701234567890"))
        self.assertEqual((product["description"], product["category"], product["unit_cost"]),
                         ("Botella", "Bebidas", 2000))
        for name in POS_BACKUP_SHEETS:
            self.assertEqual(len(tables[name]), 1)
        self.assertEqual(tables["store_orders"][0]["request_payload"], self.tables["store_orders"][0]["request_payload"])
        self.assertEqual(tables["store_sales"][0]["sold_at"], "2026-10-14T21:30:00-05:00")
        self.assertEqual(tables["store_payments"][0]["change_amount"], 2000)

    def test_json_backup_preserves_void_order_cost_null_and_ledgers(self):
        db = self.database()
        self.tables["store_products"][0]["unit_cost"] = None
        self.tables["store_sales"][0]["status"] = "VOID"
        self.tables["store_orders"][0].update(status="VOID", void_reason="Duplicada",
                                              voided_at="2026-10-15T03:00:00Z")
        with tempfile.TemporaryDirectory() as folder:
            path = db.backup(Path(folder) / "pos.json")
            tables = db.prepare_technical_backup_replacement(path)["tables"]
        self.assertIsNone(tables["store_products"][0]["unit_cost"])
        self.assertEqual(tables["store_orders"][0]["status"], "VOID")
        self.assertEqual(tables["store_orders"][0]["voided_at"], "2026-10-15T03:00:00Z")
        self.assertEqual(len(tables["store_inventory_movements"]), 1)

    def test_legacy_workbook_still_loads_without_pos_sheets_and_sku_optional(self):
        from openpyxl import load_workbook
        db = self.database()
        self.tables["store_products"][0].pop("sku")
        self.tables["store_sales"] = []
        with tempfile.TemporaryDirectory() as folder:
            path = db.export_excel(Path(folder) / "legacy.xlsx")
            book = load_workbook(path)
            for title in POS_BACKUP_SHEETS.values():
                del book[title]
            book.save(path)
            result = db.prepare_excel_replacement(path)
        self.assertNotIn("store_orders", result["tables"])
        self.assertEqual(result["tables"]["store_products"][0]["barcode"], "007701234567890")

    def test_duplicate_barcode_rejected_before_server_write(self):
        rows = {"store_products": [dict(name="Agua", barcode="0001"), dict(name="Jugo", barcode="0001")]}
        with self.assertRaisesRegex(ValueError, "código de barras repetido"):
            validate_pos_import(rows)

    def test_invisible_barcode_cannot_poison_catalog_after_import(self):
        for barcode in ("001\u200b23", "001\u00a023", "001\u008523", "001\n23"):
            with self.subTest(barcode=repr(barcode)), self.assertRaisesRegex(ValueError, "imprimibles"):
                validate_pos_import({"store_products": [dict(name="Agua", barcode=barcode)]})

    def test_import_headers_reuse_existing_aliases_and_support_new_columns(self):
        for label, expected in [("Código de barras / SKU", "sku"), ("Código", "sku"),
                                ("Nombre", "name"), ("Código de barras", "barcode"),
                                ("Stock mínimo", "low_stock_threshold"), ("Costo", "unit_cost")]:
            self.assertEqual(CloudDatabase._excel_header_key("store_products", label), expected)
        self.assertEqual(CloudDatabase._excel_cell_value("000123", "barcode"), "000123")

    def test_incomplete_pos_workbook_is_not_silently_accepted(self):
        with self.assertRaisesRegex(ValueError, "Faltan hojas"):
            validate_pos_import({"store_orders": [], "store_products": []})

    def test_void_sales_excluded_from_excel_summary_without_staff_purchase_exclusion(self):
        from openpyxl import load_workbook
        db = self.database()
        self.tables["store_sales"].append(dict(self.tables["store_sales"][0], id=6, status="VOID", total_amount=5000))
        self.tables["store_orders"][0]["customer_name"] = "STAFF"
        with tempfile.TemporaryDirectory() as folder:
            path = db.export_excel(Path(folder) / "sales.xlsx")
            book = load_workbook(path, data_only=True)
            values = {row[0]: row[1] for row in book["Resumen"].iter_rows(values_only=True)
                      if len(row) >= 2 and isinstance(row[0], str)}
            self.assertEqual(values["Ingresos por tienda"], 8000)
            self.assertEqual(values["Ingresos por membresías"], 0)
            book.close()


if __name__ == "__main__":
    unittest.main()
