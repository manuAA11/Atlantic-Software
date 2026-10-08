from __future__ import annotations
from gym_time import display_timestamp, local_datetime, timezone_for, parse_instant, local_wall_instant
from plan_forms import membership_request
from server_date import ServerDate

import json
import sqlite3
import unicodedata
from datetime import (
    date,
    datetime,
    time,
    timedelta,
    timezone,
)
from pathlib import Path
from typing import Any

from biometric import normalize_biometric_identifier
from plan_forms import validate_plan, validate_months

from cloud import CloudService
from zoneinfo import ZoneInfo
from payment_revision import (PAYMENT_EXCEL_COLUMNS, posted_income, payment_amount, validate_import_payment)  # GymSoft-PAGOS-2.0.1



TICKET_ENTRY_LIMIT = 15
TABLES = (
    "clients",
    "plans",
    "memberships",
    "checkins",
    "trainers",
    "routines",
    "exercises",
    "classes",
    "reservations",
)

CLOUD_BACKUP_TABLES = (
    *TABLES,
    "store_products",
    "store_sales",
    "accounting_expenses",
    "staff_shifts",
    "marketing_settings",
    "marketing_messages",
)

EXCEL_IMPORT_SHEETS = {
    "clients": "Clientes",
    "plans": "Planes",
    "memberships": "Membresías",
    "checkins": "Entradas",
    "store_products": "Productos",
    "store_sales": "Ventas tienda",
    "accounting_expenses": "Gastos",
    "trainers": "Personal",
    "staff_shifts": "Jornadas",
    "routines": "Rutinas",
    "exercises": "Ejercicios",
    "classes": "Clases",
    "reservations": "Reservas",
    "marketing_messages": "Mensajes marketing",
}

EXCEL_IMPORT_COLUMNS = {
    "clients": {
        "ID": "id",
        "Documento": "document",
        "Nombres": "first_name",
        "Apellidos": "last_name",
        "Teléfono": "phone",
        "Correo": "email",
        "Fecha de nacimiento": "birthdate",
        "Fecha de nacimiento alterna": "birth_date",
        "Contacto de emergencia": "emergency_contact",
        "Identificador biométrico": "biometric_identifier",
        "Autoriza WhatsApp": "whatsapp_opt_in",
        "Fecha autorización WhatsApp": "whatsapp_opt_in_at",
        "Activo": "active",
        "Creado": "created_at",
        "Actualizado": "updated_at",
    },
    "plans": {
        "ID": "id",
        "Plan": "name",
        "Duración (días)": "duration_days",
        "Vigencia (meses calendario)": "duration_months",
        "Precio": "price",
        "Límite de entradas": "entry_limit",
        "Activo": "active",
    },
    "memberships": {
        "ID": "id",
        "ID cliente": "client_id",
        "Cliente": "client_name",
        "Documento": "client_document",
        "ID plan": "plan_id",
        "Plan": "plan_name",
        "Inicio": "start_date",
        "Vencimiento": "end_date",
        "Entradas compradas": "entry_limit",
        "Entradas utilizadas antes del traslado": "initial_entries_used",
        "Valor pagado": "amount",
        "Valor pagado alterno": "amount_paid",
        "Método de pago": "payment_method",
        "Referencia": "payment_reference",
        "Notas": "notes",
        "Fecha del pago": "paid_at",
        "Registrado por": "recorded_by",
    },
    "checkins": {
        "ID": "id",
        "ID cliente": "client_id",
        "Cliente": "client_name",
        "Documento": "client_document",
        "ID membresía": "membership_id",
        "Fecha y hora": "checkin_at",
        "Método": "method",
        "Resultado": "result",
        "Registrado por": "recorded_by",
        "Notas": "notes",
    },
    "store_products": {
        "ID": "id",
        "Código de barras / SKU": "sku",
        "Producto": "name",
        "Precio de venta": "sale_price",
        "Existencias": "stock_quantity",
        "Nivel de aviso": "low_stock_threshold",
        "Tiene imagen": "has_image",
        "Activo": "active",
        "Creado": "created_at",
        "Actualizado": "updated_at",
    },
    "store_sales": {
        "ID": "id",
        "ID producto": "product_id",
        "Código": "product_sku",
        "Producto": "product_name",
        "Cantidad": "quantity",
        "Precio unitario": "unit_price",
        "Total": "total_amount",
        "Método de pago": "payment_method",
        "Referencia": "payment_reference",
        "Cliente": "customer_name",
        "Notas": "notes",
        "Registrado por": "sold_by",
        "Fecha y hora": "sold_at",
    },
    "accounting_expenses": {
        "ID": "id",
        "Fecha": "expense_date",
        "Categoría": "category",
        "Descripción": "description",
        "Proveedor / beneficiario": "vendor",
        "Valor": "amount",
        "Método de pago": "payment_method",
        "Referencia": "payment_reference",
        "Comprobante": "receipt_reference",
        "Notas": "notes",
        "Registrado por": "recorded_by",
        "Creado": "created_at",
        "Actualizado": "updated_at",
    },
    "trainers": {
        "ID": "id",
        "Nombre": "name",
        "Teléfono": "phone",
        "Correo": "email",
        "Especialidad": "specialty",
        "Pago por hora": "hourly_rate",
        "Fecha de contratación": "hire_date",
        "Notas": "notes",
        "Activo": "active",
    },
    "staff_shifts": {
        "ID": "id",
        "ID trabajador": "trainer_id",
        "Trabajador": "trainer_name",
        "Entrada": "check_in_at",
        "Salida": "check_out_at",
        "Pago por hora": "hourly_rate",
        "Notas": "notes",
        "Registrado por": "recorded_by",
    },
    "routines": {
        "ID": "id",
        "ID cliente": "client_id",
        "Cliente": "client_name",
        "ID entrenador": "trainer_id",
        "Entrenador": "trainer_name",
        "Rutina": "name",
        "Objetivo": "goal",
        "Inicio": "start_date",
        "Notas": "notes",
        "Activa": "active",
    },
    "exercises": {
        "ID": "id",
        "ID rutina": "routine_id",
        "Rutina": "routine_name",
        "Día": "day_name",
        "Ejercicio": "name",
        "Series": "sets",
        "Repeticiones": "reps",
        "Peso": "weight",
        "Notas": "notes",
        "Posición": "position",
    },
    "classes": {
        "ID": "id",
        "Clase": "name",
        "ID entrenador": "trainer_id",
        "Entrenador": "trainer_name",
        "Fecha y hora": "starts_at",
        "Cupos": "capacity",
        "Notas": "notes",
        "Estado": "status",
    },
    "reservations": {
        "ID": "id",
        "ID clase": "class_id",
        "Clase": "class_name",
        "ID cliente": "client_id",
        "Cliente": "client_name",
        "Documento": "client_document",
        "Estado": "status",
        "Registrada": "created_at",
    },
    "marketing_messages": {
        "ID": "id",
        "Automatización": "automation_type",
        "Fecha del evento": "source_date",
        "ID cliente": "client_id",
        "Cliente": "client_name",
        "WhatsApp": "phone",
        "Plantilla": "template_name",
        "Estado": "status",
        "Intentos": "attempts",
        "ID de WhatsApp": "whatsapp_message_id",
        "Detalle": "error_message",
        "Creado": "created_at",
        "Enviado": "sent_at",
    },
}
EXCEL_IMPORT_COLUMNS["memberships"].update(PAYMENT_EXCEL_COLUMNS)



class CloudDatabase:
    def marketing_workspace(self):
        from marketing_client import MarketingClient
        return MarketingClient(self).workspace()

    def __init__(self, cloud: CloudService):
        if not cloud.gym_id:
            raise ValueError("No hay un gimnasio conectado.")

        self.cloud = cloud
        self.client = cloud.client
        self.gym_id = cloud.gym_id
        self._server_date = ServerDate(getattr(cloud, 'timezone', 'America/Bogota'))

    def _execute(self, query: Any) -> Any:
        try:
            return query.execute()
        except Exception as error:
            code = str(getattr(error, "code", ""))
            message = str(error)

            if (
                code == "23505" or "duplicate key" in message.lower()
            ) and (
                "biometric" in message.lower()
                or "biométrico" in message.lower()
            ):
                raise ValueError(
                    "Este identificador biométrico ya pertenece a otro cliente."
                ) from error

            if code == "23505" or "duplicate key" in message.lower():
                raise sqlite3.IntegrityError(message) from error

            if code == "23503" or "foreign key" in message.lower():
                raise ValueError("El registro está relacionado con información inválida.") from error

            raise

    def _data(self, query: Any) -> list[dict[str, Any]]:
        response = self._execute(query)
        return list(response.data or [])
    
    def _today(self) -> date:
        return self._server_date.get(self._fetch_clock)

    def _fetch_clock(self):
        clock = self._execute(self.client.rpc('gym_local_clock', {'p_gym_id': self.gym_id})).data
        if not isinstance(clock, dict):
            raise ValueError('No se pudo consultar el reloj del gimnasio.')
        parse_instant(clock['now'])
        self.cloud.timezone = clock['timezone']
        return clock

    def invalidate_today_cache(self) -> None:
        self._server_date.invalidate()

    def _now(self) -> datetime:
        clock = self._fetch_clock()
        return parse_instant(clock['now']).astimezone(ZoneInfo(clock['timezone']))

    def _timestamp(self, value: Any) -> str:
        return display_timestamp(value, timezone_for(self))

    @staticmethod
    def _checkin_method(value: Any) -> str:
        method = str(value or "").strip()

        if method.casefold() in {
            "simulación",
            "simulacion",
            "demo",
        }:
            return "REGISTRO MANUAL"

        return method or "REGISTRO MANUAL"

    @staticmethod
    def _nested(row: dict[str, Any], key: str) -> dict[str, Any]:
        value = row.get(key) or {}
        if isinstance(value, list):
            return value[0] if value else {}
        return value

    @staticmethod
    def _json_list(value: Any) -> list[dict[str, Any]]:
        if not isinstance(value, list):
            return []

        if len(value) == 1 and isinstance(value[0], list):
            value = value[0]

        return [
            dict(item)
            for item in value
            if isinstance(item, dict)
        ]

    @staticmethod
    def _json_dict(value: Any) -> dict[str, Any]:
        if isinstance(value, list):
            value = value[0] if value else {}

        return dict(value) if isinstance(value, dict) else {}

    def seed_demo_data(self) -> None:
        # En la nube no se crean clientes de demostración.
        return

    def is_empty(self) -> bool:
        rows = self._data(
            self.client.table("clients")
            .select("id")
            .eq("gym_id", self.gym_id)
            .limit(1)
        )
        return not rows

    # Dashboard
    def dashboard_metrics(self) -> dict[str, int]:
        today_date = self._today()
        today = today_date.isoformat()

        day_start = datetime.combine(
            today_date,
            time.min,
            tzinfo=ZoneInfo(self.cloud.timezone),
        )

        next_day_start = (
            day_start + timedelta(days=1)
        )
        clients = self._data(
            self.client.table("clients")
            .select("id, active")
            .eq("gym_id", self.gym_id)
        )
        memberships = self._data(
            self.client.table("memberships")
            .select("client_id, start_date, end_date").eq("payment_status", "posted")
            .eq("gym_id", self.gym_id)
        )
        checkins = self._data(
            self.client.table("checkins")
            .select("id, checkin_at, result")
            .eq("gym_id", self.gym_id)
            .eq("result", "PERMITIDA")
            .gte(
                "checkin_at",
                day_start.isoformat(),
            )
            .lt(
                "checkin_at",
                next_day_start.isoformat(),
            )
        )
        active_ids = {
            int(item["id"])
            for item in clients
            if bool(item.get("active", True))
        }
        valid_ids = {
            int(item["client_id"])
            for item in memberships
            if item.get("start_date", "") <= today <= item.get("end_date", "")
        } & active_ids

        return {
            "active_clients": len(active_ids),
            "valid_memberships": len(valid_ids),
            "expired_memberships": max(len(active_ids) - len(valid_ids), 0),
            "checkins_today": len(checkins),
            "income_month": 0,
        }

    def session_followup(self, days=7):
        response = self._execute(self.client.rpc('session_followup', {'p_gym_id': self.gym_id, 'p_days': int(days)}))
        rows = response.data or []
        if not isinstance(rows, list):
            raise ValueError('No se pudo consultar el seguimiento de sesiones.')
        return rows

    def attendance_statistics(
        self,
        days: int = 30,
    ) -> dict[str, Any]:
        period_days = min(
            max(int(days), 7),
            90,
        )

        response = self._execute(
            self.client.rpc(
                "get_attendance_statistics",
                {
                    "p_gym_id": self.gym_id,
                    "p_days": period_days,
                },
            )
        )

        result = response.data or {}

        if isinstance(result, list):
            result = result[0] if result else {}

        if not isinstance(result, dict):
            raise ValueError(
                "Supabase devolvió estadísticas no válidas."
            )

        return result

    def membership_expiration_groups(
        self,
    ) -> dict[str, list[dict[str, Any]]]:
        today = self._today()
        groups: dict[str, list[dict[str, Any]]] = {
            "expired": [],
            "days_0_7": [],
            "days_8_15": [],
            "days_16_30": [],
            "days_31_plus": [],
        }

        clients = self._data(
            self.client.table("clients")
            .select(
                "id, document, first_name, last_name, "
                "phone, active"
            )
            .eq("gym_id", self.gym_id)
            .eq("active", True)
        )

        latest_membership: dict[int, str] = {}

        for membership in self._data(
            self.client.table("memberships")
            .select("client_id, end_date").eq("payment_status", "posted")
            .eq("gym_id", self.gym_id)
        ):
            client_id = int(membership["client_id"])
            end_date = str(
                membership.get("end_date") or ""
            )

            if end_date:
                latest_membership[client_id] = max(
                    latest_membership.get(client_id, ""),
                    end_date,
                )

        for client in clients:
            client_id = int(client["id"])
            end_text = latest_membership.get(client_id, "")
            days_remaining: int | None = None

            if end_text:
                try:
                    days_remaining = (
                        date.fromisoformat(end_text) - today
                    ).days
                except ValueError:
                    end_text = ""

            if days_remaining is None or days_remaining < 0:
                bucket = "expired"
            elif days_remaining <= 7:
                bucket = "days_0_7"
            elif days_remaining <= 15:
                bucket = "days_8_15"
            elif days_remaining <= 30:
                bucket = "days_16_30"
            else:
                bucket = "days_31_plus"

            groups[bucket].append(
                {
                    "id": client_id,
                    "document": client.get("document", ""),
                    "client_name": (
                        f"{client.get('first_name', '')} "
                        f"{client.get('last_name', '')}"
                    ).strip(),
                    "phone": client.get("phone", ""),
                    "end_date": end_text,
                    "days_remaining": days_remaining,
                }
            )

        for bucket, people in groups.items():
            people.sort(
                key=lambda person: (
                    person["days_remaining"] is None,
                    person["days_remaining"]
                    if person["days_remaining"] is not None
                    else 0,
                    person["client_name"].casefold(),
                )
            )

        return groups

    def upcoming_expirations(self, days: int = 7) -> list[dict[str, Any]]:
        today = self._today()       
        limit = today + timedelta(days=days)
        clients = {
            int(row["id"]): row
            for row in self._data(
                self.client.table("clients")
                .select("id, document, first_name, last_name, active")
                .eq("gym_id", self.gym_id)
                .eq("active", True)
            )
        }
        latest: dict[int, str] = {}

        for row in self._data(
            self.client.table("memberships")
            .select("client_id, end_date").eq("payment_status", "posted")
            .eq("gym_id", self.gym_id)
        ):
            client_id = int(row["client_id"])
            end_date = str(row["end_date"])
            latest[client_id] = max(latest.get(client_id, ""), end_date)

        result = []
        for client_id, end_text in latest.items():
            client = clients.get(client_id)
            if not client:
                continue
            end = date.fromisoformat(end_text)
            if today <= end <= limit:
                result.append(
                    {
                        "id": client_id,
                        "document": client["document"],
                        "client_name": f"{client['first_name']} {client['last_name']}",
                        "end_date": end_text,
                    }
                )

        return sorted(result, key=lambda item: item["end_date"])

    # Clientes
    def list_clients(self, search: str = "") -> list[dict[str, Any]]:
        rows = self._execute(self.client.rpc('reception_list_clients', {
            'p_gym_id': self.gym_id, 'p_search': search.strip()})).data or []
        result = []
        for row in rows:
            item = dict(row)
            item['active'] = int(bool(item.get('active', True)))
            item['birthdate'] = item.get('birthdate') or item.get('birth_date') or ''
            item['photo_path'] = item.get('photo_path') or ''
            item['membership_status'] = item.get('membership_status') or item.get('status', 'VENCIDO')
            result.append(item)
        return result

    def get_client(self, client_id: int) -> dict[str, Any] | None:
        rows = self._data(
            self.client.table("clients")
            .select("*")
            .eq("gym_id", self.gym_id)
            .eq("id", client_id)
            .limit(1)
        )
        if not rows:
            return None
        row = rows[0]
        row["active"] = int(bool(row.get("active", True)))
        row["birthdate"] = row.get("birthdate") or ""
        row["photo_path"] = row.get("photo_path") or ""
        return row

    def get_client_by_document(self, document: str) -> dict[str, Any] | None:
        rows = self._data(
            self.client.table("clients")
            .select("*")
            .eq("gym_id", self.gym_id)
            .ilike("document", document.strip())
            .limit(1)
        )
        return rows[0] if rows else None

    def get_client_by_biometric_identifier(
        self,
        identifier: str,
    ) -> dict[str, Any] | None:
        code = normalize_biometric_identifier(identifier)
        if not code:
            return None
        rows = self._data(
            self.client.table("clients")
            .select("*")
            .eq("gym_id", self.gym_id)
            .eq("biometric_identifier", code)
            .limit(1)
        )
        return rows[0] if rows else None

    def save_client(self, data: dict[str, Any], client_id: int | None = None) -> int:
        first_name = str(data.get("first_name", "")).strip()
        last_name = str(data.get("last_name", "")).strip()
        birthdate = str(data.get("birthdate", "")).strip() or None
        payload = {
            "gym_id": self.gym_id,
            "document": str(data.get("document", "")).strip(),
            "full_name": f"{first_name} {last_name}".strip(),
            "first_name": first_name,
            "last_name": last_name,
            "phone": str(data.get("phone", "")).strip(),
            "email": str(data.get("email", "")).strip(),
            "birth_date": birthdate,
            "birthdate": birthdate,
            "emergency_contact": str(data.get("emergency_contact", "")).strip(),
            "biometric_identifier": normalize_biometric_identifier(
                data.get("biometric_identifier", "")
            ),
            "photo_path": str(data.get("photo_path", "")).strip(),
            "active": bool(data.get("active", 1)),
        }

        if client_id is None:
            rows = self._data(self.client.table("clients").insert(payload))
            return int(rows[0]["id"])

        self._execute(
            self.client.table("clients")
            .update(payload)
            .eq("gym_id", self.gym_id)
            .eq("id", client_id)
        )
        return client_id

    def delete_client(self, client_id: int) -> None:
        self._execute(
            self.client.table("clients")
            .delete()
            .eq("gym_id", self.gym_id)
            .eq("id", client_id)
        )

    # Planes y mensualidades
    def ticket_followup(self):
        response = self._execute(self.client.rpc('ticket_followup', {'p_gym_id': self.gym_id}))
        rows = response.data or []
        if not isinstance(rows, list):
            raise ValueError('No se pudo consultar el saldo de las tiqueteras.')
        return rows

    def list_plans(self, active_only: bool = True) -> list[dict[str, Any]]:
        query = (
            self.client.table("plans")
            .select("*")
            .eq("gym_id", self.gym_id)
            .order("duration_days")
        )
        if active_only:
            query = query.eq("active", True)
        return self._data(query)

    def add_plan(self, name: str, duration_days: int, price: int, entry_limit: int | None = None, duration_months: int | None = None) -> int:
        clean_name, duration_days, price, entry_limit = validate_plan(name, duration_days, price, entry_limit)
        duration_months = validate_months(duration_months, entry_limit)
        payload = {
            "gym_id": self.gym_id,
            "name": clean_name,
            "duration_days": int(duration_days),
            "price": int(price),
            "entry_limit": entry_limit,
            "duration_months": duration_months,
            "active": True,
        }
        rows = self._data(self.client.table("plans").insert(payload))
        return int(rows[0]["id"])

    def update_plan(self, plan_id: int, name: str, duration_days: int, price: int, entry_limit: int | None = None, duration_months: int | None = None) -> None:
        clean_name, duration_days, price, entry_limit = validate_plan(name, duration_days, price, entry_limit)
        duration_months = validate_months(duration_months, entry_limit)
        old = self._data(self.client.table("plans").select("entry_limit").eq("gym_id", self.gym_id).eq("id", plan_id).limit(1))
        if old and (old[0].get("entry_limit") is None) != (entry_limit is None):
            used = self._data(self.client.table("memberships").select("id").eq("gym_id", self.gym_id).eq("plan_id", plan_id).limit(1))
            if used:
                raise ValueError("Este plan ya tiene compras registradas. Crea otro plan para cambiar entre días y entradas. Puedes editar el cupo de esta tiquetera para las próximas compras.")
        self._execute(
            self.client.table("plans")
            .update(
                {
                    "name": clean_name,
                    "duration_days": int(duration_days),
                    "price": int(price),
                    "entry_limit": entry_limit,
                    "duration_months": duration_months,
                }
            )
            .eq("gym_id", self.gym_id)
            .eq("id", plan_id)
        )

    def set_plan_active(self, plan_id: int, active: bool) -> None:
        self._execute(
            self.client.table("plans")
            .update({"active": bool(active)})
            .eq("gym_id", self.gym_id)
            .eq("id", plan_id)
        )


    def add_membership(
            self,
            client_id: int,
            plan_id: int,
            start_date: str,
            amount: int,
            payment_method: str,
            payment_reference: str = "",
            notes: str = "",
        carryover: bool = False,
        initial_entries_used: int = 0,
        request_id: str | None = None,
        ) -> int:
            rpc_name, params = membership_request(
                self.gym_id, client_id, plan_id, start_date, amount, payment_method,
                payment_reference, notes, carryover=carryover,
                initial_entries_used=initial_entries_used, request_id=request_id)
            response = self._execute(self.client.rpc(rpc_name, params))
            
            rows = list(response.data or [])

            if not rows:
                raise ValueError(
                    "No fue posible registrar la membresía."
                )

            return int(rows[0]["membership_id"])

    def list_memberships(self, client_id: int) -> list[dict[str, Any]]:
            rows = self._data(
                self.client.table("memberships")
                .select("*, plans(name)").eq("payment_status", "posted")
                .eq("gym_id", self.gym_id)
                .eq("client_id", client_id)
                .order("paid_at", desc=True)
                .order("id", desc=True)
            )
            result = []
            for row in rows:
                item = dict(row)
                item["plan_name"] = self._nested(row, "plans").get("name", "Sin plan")
                item["amount"] = payment_amount(item)
                item["paid_at"] = self._timestamp(item.get("paid_at"))
                result.append(item)
            return result

    def membership_snapshot(self, client_id: int) -> dict[str, Any]:
        value = self._execute(self.client.rpc("reception_membership_snapshot", {
            "p_gym_id": self.gym_id, "p_client_id": client_id})).data
        if isinstance(value, str):
            value = json.loads(value)
        if not isinstance(value, dict):
            raise ValueError("No se pudo consultar el estado de la membresía.")
        return value

    def register_checkin(self, client_id: int, method: str, override: bool = False) -> dict[str, Any]:
        result = self._execute(self.client.rpc("admin_register_checkin", {
            "p_gym_id": self.gym_id, "p_client_id": client_id,
            "p_method": method, "p_override": override,
        })).data
        if not isinstance(result, dict):
            raise ValueError("Respuesta de entrada inválida.")
        return result

    def list_checkins(self, limit: int = 100) -> list[dict[str, Any]]:
            rows = self._data(
                self.client.table("checkins")
                .select("*, clients(document, first_name, last_name)")
                .eq("gym_id", self.gym_id)
                .order("checkin_at", desc=True)
                .order("id", desc=True)
                .limit(limit)
            )
            result = []
            for row in rows:
                client = self._nested(row, "clients")
                item = dict(row)
                item["document"] = client.get("document", "")
                item["client_name"] = (
                    f"{client.get('first_name', '')} {client.get('last_name', '')}"
                ).strip()
                item["checkin_at"] = self._timestamp(item.get("checkin_at"))
                item["method"] = self._checkin_method(
                    item.get("method")
                )
                result.append(item)
            return result

    # Tienda
    def list_store_products(
        self,
        search: str = "",
        include_inactive: bool = True,
    ) -> list[dict[str, Any]]:
        response = self._execute(
            self.client.rpc(
                "admin_list_store_products",
                {
                    "p_gym_id": self.gym_id,
                    "p_search": search.strip(),
                    "p_include_inactive": bool(include_inactive),
                },
            )
        )
        return self._json_list(response.data)

    def save_store_product(
        self,
        data: dict[str, Any],
        product_id: int | None = None,
    ) -> int:
        response = self._execute(
            self.client.rpc(
                "admin_save_store_product",
                {
                    "p_gym_id": self.gym_id,
                    "p_product_id": product_id,
                    "p_name": str(data.get("name", "")).strip(),
                    "p_sku": str(data.get("sku", "")).strip(),
                    "p_sale_price": int(data.get("sale_price", 0)),
                    "p_stock_quantity": int(data.get("stock_quantity", 0)),
                    "p_low_stock_threshold": int(
                        data.get("low_stock_threshold", 5)
                    ),
                    "p_active": bool(data.get("active", True)),
                    "p_image_data": str(
                        data.get("image_data", "")
                    ),
                },
            )
        )

        value = response.data
        if isinstance(value, list):
            value = value[0] if value else None
        if isinstance(value, dict):
            value = value.get("admin_save_store_product")
        if value is None:
            raise ValueError("No fue posible guardar el producto.")
        return int(value)

    def list_store_sales(self, limit: int = 100) -> list[dict[str, Any]]:
        response = self._execute(
            self.client.rpc(
                "admin_list_store_sales",
                {
                    "p_gym_id": self.gym_id,
                    "p_limit": int(limit),
                },
            )
        )
        return self._json_list(response.data)

    def store_finance(
        self,
        date_from: str,
        date_to: str,
    ) -> dict[str, Any]:
        response = self._execute(
            self.client.rpc(
                "admin_store_finance",
                {
                    "p_gym_id": self.gym_id,
                    "p_date_from": date_from,
                    "p_date_to": date_to,
                },
            )
        )
        value = response.data
        if isinstance(value, list):
            value = value[0] if value else {}
        return dict(value) if isinstance(value, dict) else {}

    # Contabilidad y personal
    def accounting_expenses(
        self,
        date_from: str,
        date_to: str,
    ) -> dict[str, Any]:
        response = self._execute(
            self.client.rpc(
                "admin_accounting_expenses",
                {
                    "p_gym_id": self.gym_id,
                    "p_date_from": date_from,
                    "p_date_to": date_to,
                },
            )
        )
        return self._json_dict(response.data)

    def save_expense(
        self,
        data: dict[str, Any],
        expense_id: int | None = None,
    ) -> int:
        payload = {
            "gym_id": self.gym_id,
            "expense_date": str(data.get("expense_date", "")),
            "category": str(data.get("category", "")).strip(),
            "description": str(data.get("description", "")).strip(),
            "vendor": str(data.get("vendor", "")).strip(),
            "amount": int(data.get("amount", 0)),
            "payment_method": str(
                data.get("payment_method", "Efectivo")
            ).strip(),
            "payment_reference": str(
                data.get("payment_reference", "")
            ).strip(),
            "receipt_reference": str(
                data.get("receipt_reference", "")
            ).strip(),
            "notes": str(data.get("notes", "")).strip(),
            "recorded_by": self.cloud.user_id,
        }

        if expense_id is None:
            rows = self._data(
                self.client.table("accounting_expenses").insert(payload)
            )
            return int(rows[0]["id"])

        payload["updated_at"] = self._now().isoformat()
        self._execute(
            self.client.table("accounting_expenses")
            .update(payload)
            .eq("gym_id", self.gym_id)
            .eq("id", expense_id)
        )
        return expense_id

    def delete_expense(self, expense_id: int) -> None:
        self._execute(
            self.client.table("accounting_expenses")
            .delete()
            .eq("gym_id", self.gym_id)
            .eq("id", expense_id)
        )

    def staff_dashboard(
        self,
        date_from: str,
        date_to: str,
    ) -> dict[str, Any]:
        response = self._execute(
            self.client.rpc(
                "admin_staff_dashboard",
                {
                    "p_gym_id": self.gym_id,
                    "p_date_from": date_from,
                    "p_date_to": date_to,
                },
            )
        )
        return self._json_dict(response.data)

    def staff_clock(self, trainer_id: int, action: str) -> None:
        self._execute(
            self.client.rpc(
                "admin_staff_clock",
                {
                    "p_gym_id": self.gym_id,
                    "p_trainer_id": int(trainer_id),
                    "p_action": action,
                },
            )
        )

    # Entrenadores y rutinas
    def list_trainers(self, active_only: bool = False) -> list[dict[str, Any]]:
            query = (
                self.client.table("trainers")
                .select("*")
                .eq("gym_id", self.gym_id)
                .order("active", desc=True)
                .order("name")
            )
            if active_only:
                query = query.eq("active", True)

            rows = self._data(query)
            for row in rows:
                row["active"] = int(bool(row.get("active", True)))
            return rows

    def save_trainer(self, data: dict[str, Any], trainer_id: int | None = None) -> int:
            name = str(data.get("name", "")).strip()
            payload = {
                "gym_id": self.gym_id,
                "full_name": name,
                "name": name,
                "phone": str(data.get("phone", "")).strip(),
                "email": str(data.get("email", "")).strip(),
                "specialty": str(data.get("specialty", "")).strip(),
                "hourly_rate": int(data.get("hourly_rate", 0)),
                "hire_date": str(data.get("hire_date", "")).strip() or None,
                "notes": str(data.get("notes", "")).strip(),
                "created_by": self.cloud.user_id,
                "active": bool(data.get("active", 1)),
            }
            if trainer_id is None:
                rows = self._data(self.client.table("trainers").insert(payload))
                return int(rows[0]["id"])

            self._execute(
                self.client.table("trainers")
                .update(payload)
                .eq("gym_id", self.gym_id)
                .eq("id", trainer_id)
            )
            return trainer_id

    def create_routine(
            self,
            client_id: int,
            trainer_id: int | None,
            name: str,
            goal: str,
            start_date: str,
            notes: str,
        ) -> int:
            rows = self._data(
                self.client.table("routines").insert(
                    {
                        "gym_id": self.gym_id,
                        "client_id": client_id,
                        "trainer_id": trainer_id,
                        "name": name.strip(),
                        "goal": goal.strip(),
                        "start_date": start_date,
                        "notes": notes.strip(),
                        "active": True,
                    }
                )
            )
            return int(rows[0]["id"])

    def list_routines(self) -> list[dict[str, Any]]:
            rows = self._data(
                self.client.table("routines")
                .select("*, clients(first_name, last_name), trainers(name)")
                .eq("gym_id", self.gym_id)
                .order("active", desc=True)
                .order("start_date", desc=True)
                .order("id", desc=True)
            )
            exercises = self._data(
                self.client.table("exercises")
                .select("routine_id")
                .eq("gym_id", self.gym_id)
            )
            counts: dict[int, int] = {}
            for exercise in exercises:
                routine_id = int(exercise["routine_id"])
                counts[routine_id] = counts.get(routine_id, 0) + 1

            result = []
            for row in rows:
                client = self._nested(row, "clients")
                trainer = self._nested(row, "trainers")
                item = dict(row)
                item["client_name"] = (
                    f"{client.get('first_name', '')} {client.get('last_name', '')}"
                ).strip()
                item["trainer_name"] = trainer.get("name") or "Sin asignar"
                item["exercise_count"] = counts.get(int(item["id"]), 0)
                item["active"] = int(bool(item.get("active", True)))
                result.append(item)
            return result

    def add_exercise(
            self,
            routine_id: int,
            day_name: str,
            name: str,
            sets: str,
            reps: str,
            weight: str,
            notes: str,
        ) -> int:
            previous = self._data(
                self.client.table("exercises")
                .select("position")
                .eq("gym_id", self.gym_id)
                .eq("routine_id", routine_id)
                .order("position", desc=True)
                .limit(1)
            )
            position = int(previous[0].get("position", 0)) + 1 if previous else 1
            rows = self._data(
                self.client.table("exercises").insert(
                    {
                        "gym_id": self.gym_id,
                        "routine_id": routine_id,
                        "day_name": day_name.strip(),
                        "name": name.strip(),
                        "sets": sets.strip(),
                        "reps": reps.strip(),
                        "weight": weight.strip(),
                        "notes": notes.strip(),
                        "position": position,
                    }
                )
            )
            return int(rows[0]["id"])

    def list_exercises(self, routine_id: int) -> list[dict[str, Any]]:
            return self._data(
                self.client.table("exercises")
                .select("*")
                .eq("gym_id", self.gym_id)
                .eq("routine_id", routine_id)
                .order("day_name")
                .order("position")
                .order("id")
            )

        # Clases y reservas
    def create_class(
            self,
            name: str,
            trainer_id: int | None,
            starts_at: str,
            capacity: int,
            notes: str,
        ) -> int:
            rows = self._data(
                self.client.table("classes").insert(
                    {
                        "gym_id": self.gym_id,
                        "name": name.strip(),
                        "trainer_id": trainer_id,
                        "starts_at": local_wall_instant(starts_at, timezone_for(self)).isoformat(),
                        "capacity": int(capacity),
                        "notes": notes.strip(),
                        "status": "PROGRAMADA",
                    }
                )
            )
            return int(rows[0]["id"])

    def  list_classes(self) -> list[dict[str, Any]]:
            rows = self._data(
                self.client.table("classes")
                .select("*, trainers(name)")
                .eq("gym_id", self.gym_id)
                .order("starts_at")
                .order("id")
            )
            reservations = self._data(
                self.client.table("reservations")
                .select("class_id, status")
                .eq("gym_id", self.gym_id)
                .eq("status", "RESERVADA")
            )
            counts: dict[int, int] = {}
            for reservation in reservations:
                class_id = int(reservation["class_id"])
                counts[class_id] = counts.get(class_id, 0) + 1

            result = []
            for row in rows:
                item = dict(row)
                item["trainer_name"] = self._nested(row, "trainers").get("name") or "Sin asignar"
                item["reserved"] = counts.get(int(item["id"]), 0)
                item["starts_at"] = self._timestamp(item.get("starts_at"))
                result.append(item)
            return result

    def reserve_class(self, class_id: int, client_id: int) -> int:
            classes = self._data(
                self.client.table("classes")
                .select("id, capacity")
                .eq("gym_id", self.gym_id)
                .eq("id", class_id)
                .limit(1)
            )
            if not classes:
                raise ValueError("La clase no existe.")

            active = self._data(
                self.client.table("reservations")
                .select("id")
                .eq("gym_id", self.gym_id)
                .eq("class_id", class_id)
                .eq("status", "RESERVADA")
            )
            if len(active) >= int(classes[0]["capacity"]):
                raise ValueError("La clase ya no tiene cupos disponibles.")

            existing = self._data(
                self.client.table("reservations")
                .select("id, status")
                .eq("gym_id", self.gym_id)
                .eq("class_id", class_id)
                .eq("client_id", client_id)
                .limit(1)
            )
            if existing:
                if existing[0]["status"] == "RESERVADA":
                    raise ValueError("El cliente ya está reservado en esta clase.")
                reservation_id = int(existing[0]["id"])
                self._execute(
                    self.client.table("reservations")
                    .update({"status": "RESERVADA"})
                    .eq("gym_id", self.gym_id)
                    .eq("id", reservation_id)
                )
                return reservation_id

            rows = self._data(
                self.client.table("reservations").insert(
                    {
                        "gym_id": self.gym_id,
                        "class_id": class_id,
                        "client_id": client_id,
                        "status": "RESERVADA",
                    }
                )
            )
            return int(rows[0]["id"])

    def list_reservations(self, class_id: int) -> list[dict[str, Any]]:
            rows = self._data(
                self.client.table("reservations")
                .select("*, clients(document, first_name, last_name)")
                .eq("gym_id", self.gym_id)
                .eq("class_id", class_id)
            )
            result = []
            for row in rows:
                client = self._nested(row, "clients")
                item = dict(row)
                item["document"] = client.get("document", "")
                item["client_name"] = (
                    f"{client.get('first_name', '')} {client.get('last_name', '')}"
                ).strip()
                result.append(item)
            return sorted(result, key=lambda item: item["client_name"].casefold())

    def cancel_reservation(self, reservation_id: int) -> None:
            self._execute(
                self.client.table("reservations")
                .update({"status": "CANCELADA"})
                .eq("gym_id", self.gym_id)
                .eq("id", reservation_id)
            )

    # Marketing y automatizaciones de WhatsApp
    def marketing_settings(self) -> dict[str, Any]:
            defaults: dict[str, Any] = {
                "gym_id": self.gym_id,
                "whatsapp_enabled": False,
                "birthday_enabled": True,
                "expiring_enabled": True,
                "expired_enabled": True,
                "send_hour": 9,
                "warning_days": 5,
                "timezone": "America/Bogota",
                "language_code": "es_CO",
                "default_country_code": "57",
                "birthday_template": "gymsoft_cumpleanos",
                "expiring_template": "gymsoft_membresia_vence",
                "expired_template": "gymsoft_membresia_vencida",
                "birthday_extra": (
                    "Esperamos que disfrutes tu día. Pregunta por tu "
                    "beneficio de cumpleaños."
                ),
                "expiring_extra": (
                    "Renueva a tiempo para continuar entrenando sin "
                    "interrupciones."
                ),
                "expired_extra": (
                    "Te esperamos para renovar y volver a entrenar "
                    "con nosotros."
                ),
                "waba_id": "",
                "phone_number_id": "",
                "graph_api_version": "v23.0",
                "last_worker_at": "",
                "last_error": "",
            }
            rows = self._data(
                self.client.table("marketing_settings")
                .select("*")
                .eq("gym_id", self.gym_id)
                .limit(1)
            )

            if rows:
                defaults.update(rows[0])

            return defaults

    def save_marketing_settings(
        self,
        data: dict[str, Any],
    ) -> None:
            payload = {
                "gym_id": self.gym_id,
                "whatsapp_enabled": bool(
                    data.get("whatsapp_enabled", False)
                ),
                "birthday_enabled": bool(
                    data.get("birthday_enabled", True)
                ),
                "expiring_enabled": bool(
                    data.get("expiring_enabled", True)
                ),
                "expired_enabled": bool(
                    data.get("expired_enabled", True)
                ),
                "send_hour": max(
                    0,
                    min(int(data.get("send_hour", 9)), 23),
                ),
                "warning_days": max(
                    1,
                    min(int(data.get("warning_days", 5)), 30),
                ),
                "timezone": str(
                    data.get("timezone", "America/Bogota")
                ).strip(),
                "language_code": str(
                    data.get("language_code", "es_CO")
                ).strip(),
                "default_country_code": str(
                    data.get("default_country_code", "57")
                ).strip(),
                "birthday_template": str(
                    data.get("birthday_template", "")
                ).strip(),
                "expiring_template": str(
                    data.get("expiring_template", "")
                ).strip(),
                "expired_template": str(
                    data.get("expired_template", "")
                ).strip(),
                "birthday_extra": str(
                    data.get("birthday_extra", "")
                ).strip(),
                "expiring_extra": str(
                    data.get("expiring_extra", "")
                ).strip(),
                "expired_extra": str(
                    data.get("expired_extra", "")
                ).strip(),
                "waba_id": str(data.get("waba_id", "")).strip(),
                "phone_number_id": str(
                    data.get("phone_number_id", "")
                ).strip(),
                "graph_api_version": str(
                    data.get("graph_api_version", "v23.0")
                ).strip(),
                "updated_by": self.cloud.user_id,
                "updated_at": self._now().isoformat(),
            }
            existing = self._data(
                self.client.table("marketing_settings")
                .select("gym_id")
                .eq("gym_id", self.gym_id)
                .limit(1)
            )

            if existing:
                self._execute(
                    self.client.table("marketing_settings")
                    .update(payload)
                    .eq("gym_id", self.gym_id)
                )
            else:
                self._execute(
                    self.client.table("marketing_settings")
                    .insert(payload)
                )

    def marketing_contacts(
        self,
        search: str = "",
    ) -> list[dict[str, Any]]:
            clients = self._data(
                self.client.table("clients")
                .select(
                    "id, document, first_name, last_name, phone, "
                    "birthdate, active, whatsapp_opt_in, "
                    "whatsapp_opt_in_at"
                )
                .eq("gym_id", self.gym_id)
                .order("first_name")
                .order("last_name")
            )
            latest: dict[int, str] = {}

            for membership in self._data(
                self.client.table("memberships")
                .select("client_id, end_date").eq("payment_status", "posted")
                .eq("gym_id", self.gym_id)
            ):
                client_id = int(membership["client_id"])
                end_date = str(membership.get("end_date") or "")

                if end_date:
                    latest[client_id] = max(
                        latest.get(client_id, ""),
                        end_date,
                    )

            today = self._today()
            term = search.strip().casefold()
            result: list[dict[str, Any]] = []

            for client in clients:
                name = (
                    f"{client.get('first_name', '')} "
                    f"{client.get('last_name', '')}"
                ).strip()
                searchable = " ".join((
                    str(client.get("document") or ""),
                    name,
                    str(client.get("phone") or ""),
                )).casefold()

                if term and term not in searchable:
                    continue

                membership_end = latest.get(int(client["id"]), "")
                days_remaining: int | None = None

                if membership_end:
                    try:
                        days_remaining = (
                            date.fromisoformat(membership_end) - today
                        ).days
                    except ValueError:
                        membership_end = ""

                item = dict(client)
                item["client_name"] = name
                item["membership_end"] = membership_end
                item["days_remaining"] = days_remaining
                result.append(item)

            return result

    def set_marketing_opt_in(
        self,
        client_id: int,
        enabled: bool,
    ) -> None:
            now = self._now().isoformat()
            self._execute(
                self.client.table("clients")
                .update({
                    "whatsapp_opt_in": bool(enabled),
                    "whatsapp_opt_in_at": now if enabled else None,
                    "whatsapp_opt_in_by": (
                        self.cloud.user_id if enabled else None
                    ),
                })
                .eq("gym_id", self.gym_id)
                .eq("id", int(client_id))
            )

    def marketing_activity(
        self,
        limit: int = 250,
    ) -> list[dict[str, Any]]:
            return self._data(
                self.client.table("marketing_messages")
                .select("*")
                .eq("gym_id", self.gym_id)
                .order("created_at", desc=True)
                .limit(max(1, min(int(limit), 1000)))
            )

        # Respaldos de la nube
    def _cloud_table_rows(
        self,
        table: str,
        *,
        page_size: int = 750,
    ) -> list[dict[str, Any]]:
            if table == "memberships":
                value = self._execute(self.client.rpc("ticket_backup_memberships", {"p_gym_id": self.gym_id})).data
                return json.loads(value) if isinstance(value, str) else list(value or [])
            rows: list[dict[str, Any]] = []
            offset = 0

            while True:
                page = self._data(
                    self.client.table(table)
                    .select("*")
                    .eq("gym_id", self.gym_id)
                    .range(offset, offset + page_size - 1)
                )
                rows.extend(page)

                if len(page) < page_size:
                    break

                offset += page_size

            return rows

    def backup(self, destination: str | Path) -> Path:
            destination_path = Path(destination)
            destination_path.parent.mkdir(parents=True, exist_ok=True)
            payload = {
                "format": "GymSoft-CLOUD-1",
                "gym_id": self.gym_id,
                "created_at": self._now().isoformat(timespec="seconds"),
                "tables": {},
            }
            for table in CLOUD_BACKUP_TABLES:
                payload["tables"][table] = self._cloud_table_rows(table)

            destination_path.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            return destination_path

    def export_excel(self, destination: str | Path) -> Path:
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
                from openpyxl.utils import get_column_letter
            except ImportError as error:
                raise RuntimeError(
                    "Falta el componente openpyxl. Vuelve a crear los "
                    "instaladores para incluir la exportación a Excel."
                ) from error

            destination_path = Path(destination)
            destination_path.parent.mkdir(parents=True, exist_ok=True)

            table_titles = {
                "clients": "Clientes",
                "plans": "Planes",
                "memberships": "Membresías",
                "checkins": "Entradas",
                "store_products": "Productos",
                "store_sales": "Ventas tienda",
                "accounting_expenses": "Gastos",
                "trainers": "Personal",
                "staff_shifts": "Jornadas",
                "routines": "Rutinas",
                "exercises": "Ejercicios",
                "classes": "Clases",
                "reservations": "Reservas",
                "audit_logs": "Auditoría",
                "gym_users": "Usuarios",
                "marketing_settings": "Configuración marketing",
                "marketing_messages": "Mensajes marketing",
            }
            preferred_columns = {
                "clients": [
                    ("id", "ID"),
                    ("document", "Documento"),
                    ("first_name", "Nombres"),
                    ("last_name", "Apellidos"),
                    ("phone", "Teléfono"),
                    ("email", "Correo"),
                    ("birthdate", "Fecha de nacimiento"),
                    ("birth_date", "Fecha de nacimiento alterna"),
                    ("emergency_contact", "Contacto de emergencia"),
                    ("biometric_identifier", "Identificador biométrico"),
                    ("whatsapp_opt_in", "Autoriza WhatsApp"),
                    ("whatsapp_opt_in_at", "Fecha autorización WhatsApp"),
                    ("active", "Activo"),
                    ("created_at", "Creado"),
                    ("updated_at", "Actualizado"),
                ],
                "plans": [
                    ("id", "ID"),
                    ("name", "Plan"),
                    ("duration_days", "Duración (días)"),
                    ("duration_months", "Vigencia (meses calendario)"),
                    ("price", "Precio"),
                    ("entry_limit", "Límite de entradas"),
                    ("active", "Activo"),
                ],
                "memberships": [
                    ("id", "ID"),
                    ("client_id", "ID cliente"),
                    ("client_name", "Cliente"),
                    ("client_document", "Documento"),
                    ("plan_id", "ID plan"),
                    ("plan_name", "Plan"),
                    ("start_date", "Inicio"),
                    ("end_date", "Vencimiento"),
                    ("entry_limit", "Entradas compradas"),
                    ("initial_entries_used", "Entradas utilizadas antes del traslado"),
                    ("amount", "Valor pagado"),
                    ("amount_paid", "Valor pagado alterno"),
                    ("payment_method", "Método de pago"),
                    ("payment_reference", "Referencia"),
                    ("notes", "Notas"),
                    ("paid_at", "Fecha del pago"),
                    ("recorded_by", "Registrado por"),
                ] + [(key, label) for label, key in PAYMENT_EXCEL_COLUMNS.items()],
                "checkins": [
                    ("id", "ID"),
                    ("client_id", "ID cliente"),
                    ("client_name", "Cliente"),
                    ("client_document", "Documento"),
                    ("membership_id", "ID membresía"),
                    ("checkin_at", "Fecha y hora"),
                    ("method", "Método"),
                    ("result", "Resultado"),
                    ("recorded_by", "Registrado por"),
                ],
                "store_products": [
                    ("id", "ID"),
                    ("sku", "Código de barras / SKU"),
                    ("name", "Producto"),
                    ("sale_price", "Precio de venta"),
                    ("stock_quantity", "Existencias"),
                    ("low_stock_threshold", "Nivel de aviso"),
                    ("has_image", "Tiene imagen"),
                    ("active", "Activo"),
                    ("created_at", "Creado"),
                    ("updated_at", "Actualizado"),
                ],
                "store_sales": [
                    ("id", "ID"),
                    ("product_id", "ID producto"),
                    ("product_sku", "Código"),
                    ("product_name", "Producto"),
                    ("quantity", "Cantidad"),
                    ("unit_price", "Precio unitario"),
                    ("total_amount", "Total"),
                    ("payment_method", "Método de pago"),
                    ("payment_reference", "Referencia"),
                    ("customer_name", "Cliente"),
                    ("notes", "Notas"),
                    ("sold_by", "Registrado por"),
                    ("sold_at", "Fecha y hora"),
                ],
                "accounting_expenses": [
                    ("id", "ID"),
                    ("expense_date", "Fecha"),
                    ("category", "Categoría"),
                    ("description", "Descripción"),
                    ("vendor", "Proveedor / beneficiario"),
                    ("amount", "Valor"),
                    ("payment_method", "Método de pago"),
                    ("payment_reference", "Referencia"),
                    ("receipt_reference", "Comprobante"),
                    ("notes", "Notas"),
                    ("recorded_by", "Registrado por"),
                    ("created_at", "Creado"),
                    ("updated_at", "Actualizado"),
                ],
                "trainers": [
                    ("id", "ID"),
                    ("name", "Nombre"),
                    ("phone", "Teléfono"),
                    ("email", "Correo"),
                    ("specialty", "Especialidad"),
                    ("hourly_rate", "Pago por hora"),
                    ("hire_date", "Fecha de contratación"),
                    ("notes", "Notas"),
                    ("active", "Activo"),
                ],
                "staff_shifts": [
                    ("id", "ID"),
                    ("trainer_id", "ID trabajador"),
                    ("trainer_name", "Trabajador"),
                    ("check_in_at", "Entrada"),
                    ("check_out_at", "Salida"),
                    ("hourly_rate", "Pago por hora"),
                    ("notes", "Notas"),
                    ("recorded_by", "Registrado por"),
                ],
                "audit_logs": [
                    ("id", "ID"),
                    ("actor_email", "Cuenta"),
                    ("actor_role", "Rol"),
                    ("action", "Acción"),
                    ("entity_type", "Tipo de registro"),
                    ("entity_id", "ID del registro"),
                    ("summary", "Resumen"),
                    ("details", "Detalles"),
                    ("created_at", "Fecha y hora"),
                ],
                "routines": [
                    ("id", "ID"),
                    ("client_id", "ID cliente"),
                    ("client_name", "Cliente"),
                    ("trainer_id", "ID entrenador"),
                    ("trainer_name", "Entrenador"),
                    ("name", "Rutina"),
                    ("goal", "Objetivo"),
                    ("start_date", "Inicio"),
                    ("notes", "Notas"),
                    ("active", "Activa"),
                ],
                "exercises": [
                    ("id", "ID"),
                    ("routine_id", "ID rutina"),
                    ("routine_name", "Rutina"),
                    ("day_name", "Día"),
                    ("name", "Ejercicio"),
                    ("sets", "Series"),
                    ("reps", "Repeticiones"),
                    ("weight", "Peso"),
                    ("notes", "Notas"),
                    ("position", "Posición"),
                ],
                "classes": [
                    ("id", "ID"),
                    ("name", "Clase"),
                    ("trainer_id", "ID entrenador"),
                    ("trainer_name", "Entrenador"),
                    ("starts_at", "Fecha y hora"),
                    ("capacity", "Cupos"),
                    ("notes", "Notas"),
                    ("status", "Estado"),
                ],
                "reservations": [
                    ("id", "ID"),
                    ("class_id", "ID clase"),
                    ("class_name", "Clase"),
                    ("client_id", "ID cliente"),
                    ("client_name", "Cliente"),
                    ("client_document", "Documento"),
                    ("status", "Estado"),
                    ("created_at", "Registrada"),
                ],
                "marketing_messages": [
                    ("id", "ID"),
                    ("automation_type", "Automatización"),
                    ("source_date", "Fecha del evento"),
                    ("client_id", "ID cliente"),
                    ("client_name", "Cliente"),
                    ("phone", "WhatsApp"),
                    ("template_name", "Plantilla"),
                    ("status", "Estado"),
                    ("attempts", "Intentos"),
                    ("whatsapp_message_id", "ID de WhatsApp"),
                    ("error_message", "Detalle"),
                    ("created_at", "Creado"),
                    ("sent_at", "Enviado"),
                ],
            }

            raw_tables: dict[str, list[dict[str, Any]]] = {}

            for table in CLOUD_BACKUP_TABLES:
                raw_tables[table] = self._cloud_table_rows(table)

            raw_tables["audit_logs"] = sorted(
                self._cloud_table_rows("audit_logs"),
                key=lambda row: str(row.get("created_at") or ""),
                reverse=True,
            )
            raw_tables["gym_users"] = self._cloud_table_rows("gym_users")

            for product in raw_tables.get("store_products", []):
                product["has_image"] = "Sí" if product.get("image_data") else "No"
                product.pop("image_data", None)

            clients_by_id = {
                str(row.get("id")): row
                for row in raw_tables.get("clients", [])
            }
            plans_by_id = {
                str(row.get("id")): row
                for row in raw_tables.get("plans", [])
            }
            trainers_by_id = {
                str(row.get("id")): row
                for row in raw_tables.get("trainers", [])
            }
            routines_by_id = {
                str(row.get("id")): row
                for row in raw_tables.get("routines", [])
            }
            classes_by_id = {
                str(row.get("id")): row
                for row in raw_tables.get("classes", [])
            }

            def add_client_details(row: dict[str, Any]) -> None:
                client = clients_by_id.get(str(row.get("client_id")), {})
                row["client_name"] = " ".join(
                    part
                    for part in (
                        str(client.get("first_name") or "").strip(),
                        str(client.get("last_name") or "").strip(),
                    )
                    if part
                )
                row["client_document"] = str(
                    client.get("document") or ""
                )

            for membership in raw_tables.get("memberships", []):
                add_client_details(membership)
                plan = plans_by_id.get(str(membership.get("plan_id")), {})
                membership["plan_name"] = str(plan.get("name") or "")

            for checkin in raw_tables.get("checkins", []):
                add_client_details(checkin)

            for routine in raw_tables.get("routines", []):
                add_client_details(routine)
                trainer = trainers_by_id.get(
                    str(routine.get("trainer_id")),
                    {},
                )
                routine["trainer_name"] = str(trainer.get("name") or "")

            for exercise in raw_tables.get("exercises", []):
                routine = routines_by_id.get(
                    str(exercise.get("routine_id")),
                    {},
                )
                exercise["routine_name"] = str(routine.get("name") or "")

            for gym_class in raw_tables.get("classes", []):
                trainer = trainers_by_id.get(
                    str(gym_class.get("trainer_id")),
                    {},
                )
                gym_class["trainer_name"] = str(trainer.get("name") or "")

            for reservation in raw_tables.get("reservations", []):
                add_client_details(reservation)
                gym_class = classes_by_id.get(
                    str(reservation.get("class_id")),
                    {},
                )
                reservation["class_name"] = str(
                    gym_class.get("name") or ""
                )

            for shift in raw_tables.get("staff_shifts", []):
                trainer = trainers_by_id.get(
                    str(shift.get("trainer_id")),
                    {},
                )
                shift["trainer_name"] = str(trainer.get("name") or "")

            workbook = Workbook()
            summary = workbook.active
            summary.title = "Resumen"
            navy = "172235"
            dark_blue = "213D79"
            green = "16C784"
            light = "F8FAFC"
            muted_fill = "E8EEF7"
            thin = Side(style="thin", color="CBD5E1")

            summary.sheet_view.showGridLines = False
            summary.merge_cells("A1:D1")
            summary["A1"] = "GymSoft · Exportación completa"
            summary["A1"].font = Font(
                name="Segoe UI",
                size=18,
                bold=True,
                color=light,
            )
            summary["A1"].fill = PatternFill("solid", fgColor=navy)
            summary["A1"].alignment = Alignment(vertical="center")
            summary.row_dimensions[1].height = 34
            summary["A3"] = "Gimnasio"
            summary["B3"] = str(self.gym_id)
            summary["A4"] = "Generado"
            summary["B4"] = self._now().replace(
                microsecond=0
            ).isoformat(sep=" ")

            memberships = raw_tables.get("memberships", [])
            sales = raw_tables.get("store_sales", [])
            expenses = raw_tables.get("accounting_expenses", [])
            membership_income = posted_income(memberships)
            store_income = sum(
                int(row.get("total_amount") or 0)
                for row in sales
            )
            expense_total = sum(
                int(row.get("amount") or 0)
                for row in expenses
            )
            indicators = [
                ("Clientes", len(raw_tables.get("clients", [])), False),
                ("Membresías", len(memberships), False),
                ("Entradas", len(raw_tables.get("checkins", [])), False),
                ("Productos", len(raw_tables.get("store_products", [])), False),
                ("Ingresos por membresías", membership_income, True),
                ("Ingresos por tienda", store_income, True),
                ("Gastos registrados", expense_total, True),
                ("Balance histórico", membership_income + store_income - expense_total, True),
                ("Trabajadores", len(raw_tables.get("trainers", [])), False),
                ("Jornadas", len(raw_tables.get("staff_shifts", [])), False),
                (
                    "Mensajes de marketing enviados",
                    sum(
                        1
                        for row in raw_tables.get("marketing_messages", [])
                        if str(row.get("status")) == "SENT"
                    ),
                    False,
                ),
            ]
            summary["A6"] = "Indicador"
            summary["B6"] = "Valor"

            for cell in summary[6]:
                if cell.column <= 2:
                    cell.fill = PatternFill("solid", fgColor=dark_blue)
                    cell.font = Font(bold=True, color=light)

            for row_index, (label, value, is_money) in enumerate(
                indicators,
                start=7,
            ):
                summary.cell(row_index, 1, label)
                value_cell = summary.cell(row_index, 2, value)
                value_cell.number_format = (
                    '"$"#,##0'
                    if is_money
                    else '#,##0'
                )

                if row_index % 2:
                    for column in range(1, 3):
                        summary.cell(row_index, column).fill = PatternFill(
                            "solid",
                            fgColor=muted_fill,
                        )

            summary.column_dimensions["A"].width = 30
            summary.column_dimensions["B"].width = 24
            summary.freeze_panes = "A6"

            date_keys = {
                "birthdate",
                "birth_date",
                "start_date",
                "end_date",
                "expense_date",
                "hire_date",
                "source_date",
            "local_date",
            }
            datetime_keys = {
                "created_at",
                "updated_at",
                "paid_at",
                "checkin_at",
                "sold_at",
                "check_in_at",
                "check_out_at",
                "starts_at",
                "sent_at",
                "scheduled_for",
                "whatsapp_opt_in_at",
            }

            def excel_value(value: Any, key: str) -> Any:
                if value is None:
                    return ""
                if isinstance(value, bool):
                    return "Sí" if value else "No"
                if isinstance(value, (dict, list)):
                    return json.dumps(value, ensure_ascii=False)
                if key in date_keys and isinstance(value, str) and value:
                    try:
                        return date.fromisoformat(value[:10])
                    except ValueError:
                        return value
                if key in datetime_keys and isinstance(value, str) and value:
                    try:
                        parsed = datetime.fromisoformat(
                            value.replace("Z", "+00:00")
                        )
                        return local_datetime(parsed, timezone_for(self)).isoformat(sep=" ", timespec="seconds")
                    except ValueError:
                        return value
                return value

            export_order = [
                "clients",
                "plans",
                "memberships",
                "checkins",
                "store_products",
                "store_sales",
                "accounting_expenses",
                "trainers",
                "staff_shifts",
                "routines",
                "exercises",
                "classes",
                "reservations",
                "audit_logs",
                "gym_users",
                "marketing_settings",
                "marketing_messages",
            ]

            for table in export_order:
                rows = raw_tables.get(table, [])
                sheet = workbook.create_sheet(table_titles[table])
                sheet.sheet_view.showGridLines = False
                configured = preferred_columns.get(table, [])
                keys = [key for key, _label in configured]
                discovered = sorted({
                    key
                    for row in rows
                    for key in row
                    if key not in {"gym_id", "image_data"}
                })

                for key in discovered:
                    if key not in keys:
                        configured.append((key, key.replace("_", " ").title()))
                        keys.append(key)

                if not configured:
                    configured = [("sin_datos", "Sin datos")]

                last_column = get_column_letter(len(configured))
                sheet.merge_cells(f"A1:{last_column}1")
                sheet["A1"] = table_titles[table]
                sheet["A1"].fill = PatternFill("solid", fgColor=navy)
                sheet["A1"].font = Font(
                    name="Segoe UI",
                    size=15,
                    bold=True,
                    color=light,
                )
                sheet["A1"].alignment = Alignment(vertical="center")
                sheet.row_dimensions[1].height = 30

                for column, (_key, label) in enumerate(configured, start=1):
                    cell = sheet.cell(3, column, label)
                    cell.fill = PatternFill("solid", fgColor=dark_blue)
                    cell.font = Font(bold=True, color=light)
                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical="center",
                    )
                    cell.border = Border(bottom=thin)

                money_keys = {
                    "price",
                    "sale_price",
                    "amount",
                    "amount_paid",
                    "unit_price",
                    "total_amount",
                    "hourly_rate",
                    "estimated_pay",
                }

                for row_number, row in enumerate(rows, start=4):
                    for column, (key, _label) in enumerate(
                        configured,
                        start=1,
                    ):
                        cell = sheet.cell(
                            row_number,
                            column,
                            excel_value(row.get(key), key),
                        )

                        if key in money_keys:
                            cell.number_format = '"$"#,##0'
                        elif key in date_keys:
                            cell.number_format = "dd/mm/yyyy"
                        elif key in datetime_keys:
                            cell.number_format = "dd/mm/yyyy hh:mm"

                        if row_number % 2 == 0:
                            cell.fill = PatternFill(
                                "solid",
                                fgColor="F4F7FB",
                            )

                sheet.freeze_panes = "A4"

                if rows:
                    sheet.auto_filter.ref = (
                        f"A3:{last_column}{len(rows) + 3}"
                    )

                for column, (key, label) in enumerate(configured, start=1):
                    sample_lengths = [
                        len(str(row.get(key, "")))
                        for row in rows[:200]
                    ]
                    width = max([len(label), *sample_lengths], default=len(label))
                    sheet.column_dimensions[
                        get_column_letter(column)
                    ].width = min(max(width + 2, 11), 42)

            workbook.save(destination_path)
            return destination_path

    @staticmethod
    def _normalize_excel_label(value: Any) -> str:
        text = unicodedata.normalize(
            "NFKD",
            str(value or "").strip(),
        )
        text = "".join(
            character
            for character in text
            if not unicodedata.combining(character)
        )
        return " ".join(text.casefold().split())

    @staticmethod
    def _excel_reference(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @classmethod
    def _excel_header_key(
        cls,
        table: str,
        value: Any,
    ) -> str:
        normalized = cls._normalize_excel_label(value)
        aliases = {
            cls._normalize_excel_label(label): key
            for label, key in EXCEL_IMPORT_COLUMNS.get(table, {}).items()
        }

        if normalized in aliases:
            return aliases[normalized]

        raw_key = normalized.replace(" / ", "_")
        raw_key = "".join(
            character if character.isalnum() else "_"
            for character in raw_key
        )
        return "_".join(part for part in raw_key.split("_") if part)

    @classmethod
    def _excel_cell_value(
        cls,
        value: Any,
        key: str,
    ) -> Any:
        if value is None:
            return None
        if key in ("ticket_consumption_archive", "freeze_archive"):
            parsed = json.loads(value) if isinstance(value, str) else value
            if not isinstance(parsed, list):
                raise ValueError("El historial de consumo de la tiquetera no es válido.")
            return parsed

        date_fields = {
            "birthdate",
            "birth_date",
            "start_date",
            "end_date",
            "expense_date",
            "hire_date",
            "source_date",
            "local_date",
        }
        datetime_fields = {
            "created_at",
            "updated_at",
            "paid_at",
            "checkin_at",
            "sold_at",
            "check_in_at",
            "check_out_at",
            "starts_at",
            "sent_at",
            "scheduled_for",
            "whatsapp_opt_in_at",
        }
        boolean_fields = {
            "ticket_consumed", "already_consumed_today",
            "active",
            "whatsapp_opt_in",
        }
        integer_fields = {
            "duration_days",
            "price",
            "entry_limit",
            "duration_months",
            "initial_entries_used",
            "amount",
            "amount_paid",
            "quantity",
            "unit_price",
            "total_amount",
            "stock_quantity",
            "low_stock_threshold",
            "hourly_rate",
            "capacity",
            "position",
            "sort_order",
            "attempts",
        }
        identifier_fields = {
            "id",
            "client_id",
            "plan_id",
            "membership_id",
            "product_id",
            "trainer_id",
            "routine_id",
            "class_id",
            "document",
            "phone",
            "sku",
            "biometric_identifier",
            "payment_reference",
            "receipt_reference",
            "whatsapp_message_id",
            "recorded_by",
            "sold_by",
        }

        if isinstance(value, datetime):
            if key in date_fields:
                return value.date().isoformat()
            return parse_instant(value).isoformat() if key in datetime_fields else value.isoformat()

        if isinstance(value, date):
            if key in datetime_fields:
                raise ValueError('Una fecha sin hora no puede importarse como evento. Revisa el registro original.')
            return value.isoformat()

        if key in identifier_fields:
            return cls._excel_reference(value)

        if key in boolean_fields:
            if isinstance(value, bool):
                return value
            normalized = cls._normalize_excel_label(value)
            if normalized in {"si", "sí", "true", "verdadero", "1", "activo"}:
                return True
            if normalized in {"no", "false", "falso", "0", "inactivo"}:
                return False
            raise ValueError(
                f"El valor '{value}' no representa Sí o No."
            )

        if key in integer_fields:
            if isinstance(value, bool):
                return int(value)
            if isinstance(value, (int, float)):
                return int(round(float(value)))

            text = str(value).strip().replace("$", "").replace(" ", "")

            if not text:
                return None

            if "." in text and "," in text:
                if text.rfind(",") > text.rfind("."):
                    text = text.replace(".", "").replace(",", ".")
                else:
                    text = text.replace(",", "")
            elif text.count(".") >= 1:
                groups = text.split(".")
                if all(len(group) == 3 for group in groups[1:]):
                    text = "".join(groups)
            elif text.count(",") >= 1:
                groups = text.split(",")
                if all(len(group) == 3 for group in groups[1:]):
                    text = "".join(groups)
                else:
                    text = text.replace(",", ".")

            return int(round(float(text)))

        if key in date_fields or key in datetime_fields:
            text = str(value).strip()

            if not text:
                return None

            try:
                parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
            except ValueError:
                parsed = None

                for pattern in (
                    "%d/%m/%Y %H:%M",
                    "%d/%m/%Y",
                    "%d-%m-%Y %H:%M",
                    "%d-%m-%Y",
                    "%Y-%m-%d %H:%M",
                    "%Y-%m-%d",
                ):
                    try:
                        parsed = datetime.strptime(text, pattern)
                        break
                    except ValueError:
                        continue

            if parsed is None:
                raise ValueError(
                    f"La fecha '{value}' no tiene un formato válido."
                )

            if key in date_fields:
                return parsed.date().isoformat()
            return parse_instant(parsed).isoformat()

        if isinstance(value, float) and value.is_integer():
            return int(value)

        if isinstance(value, str):
            return value.strip()

        return value

    def prepare_excel_replacement(
        self,
        source: str | Path,
    ) -> dict[str, Any]:
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise RuntimeError(
                "Falta el componente openpyxl. Vuelve a crear los "
                "instaladores para habilitar la importación de Excel."
            ) from error

        source_path = Path(source)

        if not source_path.is_file():
            raise ValueError("No se encontró el archivo de Excel.")

        if source_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
            raise ValueError(
                "Selecciona un archivo .xlsx creado o preparado para GymSoft."
            )

        try:
            workbook = load_workbook(
                source_path,
                read_only=True,
                data_only=True,
                keep_links=False,
            )
        except Exception as error:
            raise ValueError(
                "No fue posible abrir el Excel. Verifica que no esté dañado."
            ) from error

        missing_sheets = [
            sheet_name
            for sheet_name in EXCEL_IMPORT_SHEETS.values()
            if sheet_name not in workbook.sheetnames
        ]

        if missing_sheets:
            workbook.close()
            raise ValueError(
                "El Excel no tiene el formato completo de GymSoft. "
                "Faltan estas hojas:\n- "
                + "\n- ".join(missing_sheets)
            )

        tables: dict[str, list[dict[str, Any]]] = {}
        errors: list[str] = []

        try:
            for table, sheet_name in EXCEL_IMPORT_SHEETS.items():
                sheet = workbook[sheet_name]
                sheet_values = list(sheet.iter_rows(values_only=True))
                alias_keys = set(EXCEL_IMPORT_COLUMNS.get(table, {}).values())
                best_row = 0
                best_headers: list[str] = []
                best_score = -1

                for row_number, values in enumerate(
                    sheet_values[:10],
                    start=1,
                ):
                    headers = [
                        self._excel_header_key(table, value)
                        if value not in (None, "")
                        else ""
                        for value in values
                    ]
                    score = sum(
                        1 for header in headers if header in alias_keys
                    )

                    if score > best_score:
                        best_row = row_number
                        best_headers = headers
                        best_score = score

                if best_score <= 0:
                    errors.append(
                        f"{sheet_name}: no se encontraron encabezados válidos."
                    )
                    tables[table] = []
                    continue

                duplicated = {
                    header
                    for header in best_headers
                    if header and best_headers.count(header) > 1
                }

                if duplicated:
                    errors.append(
                        f"{sheet_name}: hay columnas repetidas ({', '.join(sorted(duplicated))})."
                    )

                rows: list[dict[str, Any]] = []

                for row_number, values in enumerate(
                    sheet_values[best_row:],
                    start=best_row + 1,
                ):
                    row: dict[str, Any] = {}
                    has_value = False

                    for column, key in enumerate(best_headers, start=1):
                        if not key or key == "sin_datos":
                            continue

                        value = (
                            values[column - 1]
                            if column <= len(values)
                            else None
                        )

                        if value is None or (
                            isinstance(value, str) and not value.strip()
                        ):
                            continue

                        has_value = True

                        try:
                            cleaned = self._excel_cell_value(value, key)
                        except (TypeError, ValueError) as error:
                            errors.append(
                                f"{sheet_name}, fila {row_number}, "
                                f"columna {column}: {error}"
                            )
                            continue

                        if cleaned is not None and cleaned != "":
                            row[key] = cleaned

                    if has_value:
                        row["__excel_row"] = row_number
                        row["__excel_table"] = table
                        rows.append(row)

                tables[table] = rows
        finally:
            workbook.close()

        if errors:
            raise ValueError(
                "El Excel contiene errores:\n\n- "
                + "\n- ".join(errors[:20])
                + (
                    f"\n\nY {len(errors) - 20} errores adicionales."
                    if len(errors) > 20
                    else ""
                )
            )

        validation_errors: list[str] = []
        identifiers: dict[str, set[str]] = {}

        for table, rows in tables.items():
            identifiers[table] = set()

            for index, row in enumerate(rows, start=1):
                identifier = self._excel_reference(row.get("id"))

                if not identifier:
                    identifier = f"__{table}_{index}"
                    row["id"] = identifier

                if identifier in identifiers[table]:
                    validation_errors.append(
                        f"{EXCEL_IMPORT_SHEETS[table]}, fila "
                        f"{row['__excel_row']}: ID repetido ({identifier})."
                    )
                identifiers[table].add(identifier)

        if not tables["clients"]:
            validation_errors.append(
                "La hoja Clientes debe contener al menos una persona."
            )

        if not tables["plans"]:
            validation_errors.append(
                "La hoja Planes debe contener al menos un plan."
            )

        def normalized_key(value: Any) -> str:
            return self._normalize_excel_label(value)

        def unique_map(
            table: str,
            field: str,
            label: str,
            *,
            required: bool = False,
            duplicates_are_errors: bool = True,
        ) -> dict[str, str]:
            result: dict[str, str] = {}
            ambiguous: set[str] = set()

            for row in tables[table]:
                value = normalized_key(row.get(field))

                if not value:
                    if required:
                        validation_errors.append(
                            f"{EXCEL_IMPORT_SHEETS[table]}, fila "
                            f"{row['__excel_row']}: falta {label}."
                        )
                    continue

                if value in result or value in ambiguous:
                    if duplicates_are_errors:
                        validation_errors.append(
                            f"{EXCEL_IMPORT_SHEETS[table]}, fila "
                            f"{row['__excel_row']}: {label} repetido "
                            f"({row.get(field)})."
                        )
                    else:
                        result.pop(value, None)
                        ambiguous.add(value)
                    continue

                result[value] = self._excel_reference(row["id"])

            return result

        clients_by_document = unique_map(
            "clients",
            "document",
            "el documento",
            required=True,
        )
        plans_by_name = unique_map(
            "plans",
            "name",
            "el nombre del plan",
            required=True,
        )
        products_by_sku = unique_map(
            "store_products",
            "sku",
            "el código del producto",
            required=True,
        )
        trainers_by_name = unique_map(
            "trainers",
            "name",
            "el nombre del trabajador",
            required=True,
            duplicates_are_errors=False,
        )
        routines_by_name = unique_map(
            "routines",
            "name",
            "el nombre de la rutina",
            duplicates_are_errors=False,
        )
        classes_by_name = unique_map(
            "classes",
            "name",
            "el nombre de la clase",
            duplicates_are_errors=False,
        )

        biometric_codes: set[str] = set()

        for row in tables["clients"]:
            if not str(row.get("first_name") or "").strip():
                validation_errors.append(
                    f"Clientes, fila {row['__excel_row']}: faltan los nombres."
                )

            biometric = normalize_biometric_identifier(
                row.get("biometric_identifier")
            )

            if biometric:
                if biometric in biometric_codes:
                    validation_errors.append(
                        f"Clientes, fila {row['__excel_row']}: identificador "
                        f"biométrico repetido ({biometric})."
                    )
                biometric_codes.add(biometric)
                row["biometric_identifier"] = biometric

        def resolve_reference(
            row: dict[str, Any],
            field: str,
            target_table: str,
            alternatives: tuple[tuple[str, dict[str, str]], ...] = (),
            *,
            required: bool,
        ) -> None:
            reference = self._excel_reference(row.get(field))

            if reference in identifiers[target_table]:
                row[field] = reference
                return

            for alternative_field, mapping in alternatives:
                alternative = normalized_key(row.get(alternative_field))

                if alternative and alternative in mapping:
                    row[field] = mapping[alternative]
                    return

            if reference:
                validation_errors.append(
                    f"{EXCEL_IMPORT_SHEETS[row['__excel_table']]}, "
                    f"fila {row['__excel_row']}: {field} apunta a un registro inexistente."
                )
            elif required:
                validation_errors.append(
                    f"{EXCEL_IMPORT_SHEETS[row['__excel_table']]}, fila "
                    f"{row['__excel_row']}: falta la relación {field}."
                )
            else:
                row.pop(field, None)

        for row in tables["memberships"]:
            validate_import_payment(row)
            resolve_reference(
                row,
                "client_id",
                "clients",
                (("client_document", clients_by_document),),
                required=True,
            )
            resolve_reference(
                row,
                "plan_id",
                "plans",
                (("plan_name", plans_by_name),),
                required=True,
            )

            if not row.get("start_date") or not row.get("end_date"):
                validation_errors.append(
                    f"Membresías, fila {row['__excel_row']}: faltan inicio o vencimiento."
                )
            elif str(row["end_date"]) < str(row["start_date"]):
                validation_errors.append(
                    f"Membresías, fila {row['__excel_row']}: el vencimiento "
                    "es anterior al inicio."
                )

            if "amount" not in row and "amount_paid" in row:
                row["amount"] = row["amount_paid"]
            if "amount_paid" not in row and "amount" in row:
                row["amount_paid"] = row["amount"]
            row.setdefault("amount", 0)
            row.setdefault("amount_paid", row["amount"])
            row.setdefault("payment_method", "Efectivo")

        for row in tables["checkins"]:
            resolve_reference(
                row,
                "client_id",
                "clients",
                (("client_document", clients_by_document),),
                required=True,
            )
            resolve_reference(
                row,
                "membership_id",
                "memberships",
                required=False,
            )

        for row in tables["store_sales"]:
            resolve_reference(
                row,
                "product_id",
                "store_products",
                (("product_sku", products_by_sku),),
                required=True,
            )

        for row in tables["staff_shifts"]:
            resolve_reference(
                row,
                "trainer_id",
                "trainers",
                (("trainer_name", trainers_by_name),),
                required=True,
            )

        for row in tables["routines"]:
            resolve_reference(
                row,
                "client_id",
                "clients",
                (("client_name", {}),),
                required=True,
            )
            resolve_reference(
                row,
                "trainer_id",
                "trainers",
                (("trainer_name", trainers_by_name),),
                required=False,
            )

        for row in tables["exercises"]:
            resolve_reference(
                row,
                "routine_id",
                "routines",
                (("routine_name", routines_by_name),),
                required=True,
            )

        for row in tables["classes"]:
            resolve_reference(
                row,
                "trainer_id",
                "trainers",
                (("trainer_name", trainers_by_name),),
                required=False,
            )

        for row in tables["reservations"]:
            resolve_reference(
                row,
                "class_id",
                "classes",
                (("class_name", classes_by_name),),
                required=True,
            )
            resolve_reference(
                row,
                "client_id",
                "clients",
                (("client_document", clients_by_document),),
                required=True,
            )

        for row in tables["marketing_messages"]:
            resolve_reference(
                row,
                "client_id",
                "clients",
                required=False,
            )

        if validation_errors:
            raise ValueError(
                "El Excel no puede importarse todavía:\n\n- "
                + "\n- ".join(validation_errors[:25])
                + (
                    f"\n\nY {len(validation_errors) - 25} errores adicionales."
                    if len(validation_errors) > 25
                    else ""
                )
            )

        clean_tables: dict[str, list[dict[str, Any]]] = {}

        for table, rows in tables.items():
            clean_tables[table] = [
                {
                    key: value
                    for key, value in row.items()
                    if not key.startswith("__excel_")
                }
                for row in rows
            ]

        counts = {
            table: len(rows)
            for table, rows in clean_tables.items()
        }

        return {
            "format": "GymSoft-EXCEL-REPLACE-2",
            "source_name": source_path.name,
            "tables": clean_tables,
            "counts": counts,
            "total_rows": sum(counts.values()),
        }

    def prepare_technical_backup_replacement(
        self,
        source: str | Path,
    ) -> dict[str, Any]:
        source_path = Path(source)

        if not source_path.is_file():
            raise ValueError("No se encontró el respaldo técnico.")

        try:
            payload = json.loads(source_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as error:
            raise ValueError(
                "El archivo seleccionado no es un respaldo JSON válido."
            ) from error

        if payload.get("format") != "GymSoft-CLOUD-1":
            raise ValueError(
                "El archivo no fue creado por el sistema de respaldos de GymSoft."
            )

        source_tables = payload.get("tables")

        if not isinstance(source_tables, dict):
            raise ValueError("El respaldo no contiene las tablas esperadas.")

        tables: dict[str, list[dict[str, Any]]] = {}

        for table in EXCEL_IMPORT_SHEETS:
            rows = source_tables.get(table, [])

            if not isinstance(rows, list) or any(
                not isinstance(row, dict) for row in rows
            ):
                raise ValueError(
                    f"La tabla {table} del respaldo está dañada."
                )

            tables[table] = [dict(row) for row in rows]

        if not tables["clients"] or not tables["plans"]:
            raise ValueError(
                "El respaldo debe contener clientes y al menos un plan."
            )

        for table, rows in tables.items():
            for index, row in enumerate(rows, start=1):
                if not self._excel_reference(row.get("id")):
                    row["id"] = f"__backup_{table}_{index}"

        counts = {
            table: len(rows)
            for table, rows in tables.items()
        }

        return {
            "format": "GymSoft-EXCEL-REPLACE-2",
            "source_name": source_path.name,
            "tables": tables,
            "counts": counts,
            "total_rows": sum(counts.values()),
        }

    def replace_from_excel(
        self,
        prepared: dict[str, Any],
    ) -> dict[str, Any]:
        if prepared.get("format") != "GymSoft-EXCEL-REPLACE-2":
            raise ValueError("La importación preparada no es válida.")

        response = self._execute(
            self.client.rpc(
                "admin_replace_gym_from_excel",
                {
                    "p_data": {
                        "format": prepared["format"],
                        "tables": prepared["tables"],
                    }
                },
            )
        )
        result: Any = response.data

        if isinstance(result, list):
            result = result[0] if result else {}

        if not isinstance(result, dict):
            raise RuntimeError(
                "Supabase no confirmó el resultado de la importación."
            )

        return result

    def restore(self, source: str | Path) -> Path:
        source_path = Path(source)
        prepared = self.prepare_technical_backup_replacement(source_path)
        automatic_backup = source_path.parent / (
            f"GymSoft-antes-importacion-{datetime.now():%Y%m%d-%H%M%S}.json"
        )
        self.backup(automatic_backup)
        # One server transaction validates and remaps every relationship. A
        # failure rolls back the entire operation, including daily balances.
        self.replace_from_excel(prepared)
        return automatic_backup

    def _delete_cloud_data(self, include_plans: bool = False) -> None:
            order = (
                "marketing_messages",
                "staff_shifts",
                "store_sales",
                "accounting_expenses",
                "reservations",
                "exercises",
                "checkins",
                "classes",
                "routines",
                "memberships",
                "trainers",
                "clients",
                "store_products",
                "marketing_settings",
            )
            for table in order:
                self._execute(
                    self.client.table(table)
                    .delete()
                    .eq("gym_id", self.gym_id)
                )
            if include_plans:
                self._execute(
                    self.client.table("plans")
                    .delete()
                    .eq("gym_id", self.gym_id)
                )

    def _restore_rows(self, tables: dict[str, list[dict[str, Any]]]) -> None:
            order = (
                "marketing_settings",
                "clients",
                "plans",
                "store_products",
                "trainers",
                "accounting_expenses",
                "memberships",
                "routines",
                "classes",
                "checkins",
                "exercises",
                "reservations",
                "store_sales",
                "staff_shifts",
                "marketing_messages",
            )
            for table in order:
                rows = []
                for original in tables.get(table, []):
                    row = dict(original)
                    row["gym_id"] = self.gym_id
                    rows.append(row)
                if rows:
                    self._execute(self.client.table(table).insert(rows))

        # Importación única de SQLite
    def import_sqlite(self, source: str | Path) -> dict[str, int]:
            source_path = Path(source)
            if not source_path.is_file():
                raise ValueError("No se encontró la base de datos local.")

            connection = sqlite3.connect(source_path)
            connection.row_factory = sqlite3.Row
            try:
                existing = {
                    row[0]
                    for row in connection.execute(
                        "SELECT name FROM sqlite_master WHERE type = 'table'"
                    ).fetchall()
                }
                required = set(TABLES)
                if not required.issubset(existing):
                    raise ValueError("El archivo no es una base válida de GymSoft.")

                local = {
                    table: [dict(row) for row in connection.execute(f"SELECT * FROM {table}")]
                    for table in TABLES
                }
            finally:
                connection.close()

            return self._import_local_rows(local)

    def _import_local_rows(
            self,
            local: dict[str, list[dict[str, Any]]],
        ) -> dict[str, int]:
            counts = {table: 0 for table in TABLES}

            # Planes: se relacionan por nombre para no duplicar los tres iniciales.
            cloud_plans = self.list_plans(active_only=False)
            plans_by_name = {
                str(plan["name"]).strip().casefold(): plan
                for plan in cloud_plans
            }
            plan_map: dict[int, int] = {}

            for plan in local["plans"]:
                name = str(plan["name"]).strip()
                existing = plans_by_name.get(name.casefold())
                if existing:
                    cloud_id = int(existing["id"])
                    self.update_plan(
                        cloud_id,
                        name,
                        int(plan["duration_days"]),
                        int(plan["price"]),
                    )
                    self.set_plan_active(cloud_id, bool(plan.get("active", 1)))
                else:
                    cloud_id = self.add_plan(
                        name,
                        int(plan["duration_days"]),
                        int(plan["price"]),
                    )
                    self.set_plan_active(cloud_id, bool(plan.get("active", 1)))
                    plans_by_name[name.casefold()] = {"id": cloud_id, "name": name}

                plan_map[int(plan["id"])] = cloud_id
                counts["plans"] += 1

            client_map: dict[int, int] = {}
            existing_clients = {
                str(row["document"]).strip().casefold(): row
                for row in self.list_clients()
            }
            for client in local["clients"]:
                data = {
                    "document": client.get("document", ""),
                    "first_name": client.get("first_name", ""),
                    "last_name": client.get("last_name", ""),
                    "phone": client.get("phone", ""),
                    "email": client.get("email", ""),
                    "birthdate": client.get("birthdate", ""),
                    "emergency_contact": client.get("emergency_contact", ""),
                    "biometric_identifier": client.get(
                        "biometric_identifier",
                        "",
                    ),
                    "photo_path": client.get("photo_path", ""),
                    "active": client.get("active", 1),
                }
                key = str(data["document"]).strip().casefold()
                existing = existing_clients.get(key)
                if existing:
                    cloud_id = self.save_client(data, int(existing["id"]))
                else:
                    cloud_id = self.save_client(data)
                client_map[int(client["id"])] = cloud_id
                counts["clients"] += 1

            trainer_map: dict[int, int] = {}
            for trainer in local["trainers"]:
                cloud_id = self.save_trainer(
                    {
                        "name": trainer.get("name", ""),
                        "phone": trainer.get("phone", ""),
                        "email": trainer.get("email", ""),
                        "specialty": trainer.get("specialty", ""),
                        "active": trainer.get("active", 1),
                    }
                )
                trainer_map[int(trainer["id"])] = cloud_id
                counts["trainers"] += 1

            membership_map: dict[int, int] = {}
            plans_now = {
                int(row["id"]): row
                for row in self.list_plans(active_only=False)
            }
            existing_memberships = self._data(
                self.client.table("memberships")
                .select(
                    "id, client_id, plan_id, start_date, end_date, "
                    "amount, payment_method"
                )
                .eq("gym_id", self.gym_id)
            )
            existing_membership_by_key = {
                (
                    int(row["client_id"]),
                    int(row["plan_id"]),
                    str(row.get("start_date") or ""),
                    str(row.get("end_date") or ""),
                    int(row.get("amount") or 0),
                    str(row.get("payment_method") or "").strip().casefold(),
                ): int(row["id"])
                for row in existing_memberships
            }
            for membership in local["memberships"]:
                cloud_plan_id = plan_map[int(membership["plan_id"])]
                plan = plans_now[cloud_plan_id]
                amount = int(membership.get("amount", 0))
                cloud_client_id = client_map[int(membership["client_id"])]
                payment_method = str(
                    membership.get("payment_method") or "Efectivo"
                ).strip()
                membership_key = (
                    cloud_client_id,
                    cloud_plan_id,
                    str(membership["start_date"]),
                    str(membership["end_date"]),
                    amount,
                    payment_method.casefold(),
                )
                existing_membership_id = existing_membership_by_key.get(
                    membership_key
                )
                if existing_membership_id is not None:
                    membership_map[int(membership["id"])] = (
                        existing_membership_id
                    )
                    continue

                payload = {
                    "gym_id": self.gym_id,
                    "client_id": cloud_client_id,
                    "plan_id": cloud_plan_id,
                    "start_date": membership["start_date"],
                    "end_date": membership["end_date"],
                    "entry_limit": plan.get("entry_limit"),
                    "amount": amount,
                    "amount_paid": amount,
                    "payment_method": payment_method,
                    "payment_reference": membership.get("payment_reference", ""),
                    "paid_at": membership.get("paid_at") or self._now().isoformat(),
                    "notes": membership.get("notes", ""),
                }
                rows = self._data(self.client.table("memberships").insert(payload))
                new_membership_id = int(rows[0]["id"])
                membership_map[int(membership["id"])] = new_membership_id
                existing_membership_by_key[membership_key] = new_membership_id
                counts["memberships"] += 1

            routine_map: dict[int, int] = {}
            for routine in local["routines"]:
                local_trainer_id = routine.get("trainer_id")
                cloud_id = self.create_routine(
                    client_id=client_map[int(routine["client_id"])],
                    trainer_id=(
                        trainer_map.get(int(local_trainer_id))
                        if local_trainer_id is not None
                        else None
                    ),
                    name=str(routine.get("name", "")),
                    goal=str(routine.get("goal", "")),
                    start_date=str(routine.get("start_date") or self._today().isoformat()),
                    notes=str(routine.get("notes", "")),
                )
                routine_map[int(routine["id"])] = cloud_id
                counts["routines"] += 1

            class_map: dict[int, int] = {}
            for gym_class in local["classes"]:
                local_trainer_id = gym_class.get("trainer_id")
                cloud_id = self.create_class(
                    name=str(gym_class.get("name", "")),
                    trainer_id=(
                        trainer_map.get(int(local_trainer_id))
                        if local_trainer_id is not None
                        else None
                    ),
                    starts_at=str(gym_class["starts_at"]),
                    capacity=int(gym_class["capacity"]),
                    notes=str(gym_class.get("notes", "")),
                )
                if gym_class.get("status") and gym_class["status"] != "PROGRAMADA":
                    self._execute(
                        self.client.table("classes")
                        .update({"status": gym_class["status"]})
                        .eq("gym_id", self.gym_id)
                        .eq("id", cloud_id)
                    )
                class_map[int(gym_class["id"])] = cloud_id
                counts["classes"] += 1

            for checkin in local["checkins"]:
                local_membership_id = checkin.get("membership_id")
                self._execute(
                    self.client.table("checkins").insert(
                        {
                            "gym_id": self.gym_id,
                            "client_id": client_map[int(checkin["client_id"])],
                            "membership_id": (
                                membership_map.get(int(local_membership_id))
                                if local_membership_id is not None
                                else None
                            ),
                            "checkin_at": checkin.get("checkin_at") or self._now().isoformat(),
                            "method": self._checkin_method(
                                checkin.get("method")
                            ),
                            "result": checkin.get("result", "DENEGADA"),
                            "notes": checkin.get("notes", ""),
                        }
                    )
                )
                counts["checkins"] += 1

            for exercise in local["exercises"]:
                position = int(exercise.get("position", 0))
                self._execute(
                    self.client.table("exercises").insert(
                        {
                            "gym_id": self.gym_id,
                            "routine_id": routine_map[int(exercise["routine_id"])],
                            "day_name": exercise.get("day_name", ""),
                            "name": exercise.get("name", ""),
                            "sets": str(exercise.get("sets", "")),
                            "reps": str(exercise.get("reps", "")),
                            "weight": str(exercise.get("weight", "")),
                            "notes": exercise.get("notes", ""),
                            "position": position,
                        }
                    )
                )
                counts["exercises"] += 1

            for reservation in local["reservations"]:
                self._execute(
                    self.client.table("reservations").insert(
                        {
                            "gym_id": self.gym_id,
                            "class_id": class_map[int(reservation["class_id"])],
                            "client_id": client_map[int(reservation["client_id"])],
                            "status": reservation.get("status", "RESERVADA"),
                            "created_at": reservation.get("created_at") or self._now().isoformat(),
                        }
                    )
                )
                counts["reservations"] += 1

            return counts
